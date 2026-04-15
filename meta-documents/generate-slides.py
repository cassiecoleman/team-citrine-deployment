#!/usr/bin/env python3
"""Generate the P3 Backend Demo slide deck for Team Citrine."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

DARK_BG = RGBColor(0x1A, 0x1A, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x3B, 0x82, 0xF6)  # blue-500
GREEN = RGBColor(0x22, 0xC5, 0x5E)
GRAY = RGBColor(0x94, 0xA3, 0xB8)
RED = RGBColor(0xEF, 0x44, 0x44)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_bg(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BG


def add_text(slide, left, top, width, height, text, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = align
    return tf


def add_bullet(tf, text, size=16, color=WHITE, level=0):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.level = level
    return p


# ─── Slide 1: Title ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 1, 1.5, 11, 1.5, "Ultra — P3 Backend Demo", size=44, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, 1, 3.2, 11, 0.8, "Team Citrine", size=28, color=ACCENT, align=PP_ALIGN.CENTER)
add_text(slide, 1, 4.2, 11, 0.6, "Jacob Moore  |  Cassie Coleman  |  Derron Dowdy", size=20, color=GRAY, align=PP_ALIGN.CENTER)
add_text(slide, 1, 5.2, 11, 0.6, "Spring 2026 — AI Tools for Software Development", size=16, color=GRAY, align=PP_ALIGN.CENTER)

# ─── Slide 2: Tech Stack ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.5, 11, 0.8, "Tech Stack", size=36, bold=True)
tf = add_text(slide, 0.8, 1.5, 5.5, 5, "Frontend", size=24, bold=True, color=ACCENT)
add_bullet(tf, "Next.js 16 + React 19 + TypeScript")
add_bullet(tf, "Tailwind CSS v4")
add_bullet(tf, "Leaflet maps + OpenStreetMap")
add_bullet(tf, "Supabase Realtime (WebSocket)")
tf2 = add_text(slide, 7, 1.5, 5.5, 5, "Backend", size=24, bold=True, color=GREEN)
add_bullet(tf2, "Supabase (hosted PostgreSQL)")
add_bullet(tf2, "Row-Level Security policies")
add_bullet(tf2, "Next.js Server Actions")
add_bullet(tf2, "Stripe (test mode) for payments")
add_bullet(tf2, "Zod validation on all inputs")
add_bullet(tf2, "Vitest + Playwright for testing")

# ─── Slide 3: Architecture ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.5, 11, 0.8, "Architecture Overview", size=36, bold=True)
tf = add_text(slide, 0.8, 1.6, 11, 5, "", size=18)
add_bullet(tf, "3 user roles: Rider, Driver, Admin — each with dedicated route groups")
add_bullet(tf, "Session-based auth via Supabase cookies — middleware enforces role access")
add_bullet(tf, "Server Actions as the API layer — no separate REST endpoints needed")
add_bullet(tf, "Service Role Client bypasses RLS for server-side operations")
add_bullet(tf, "Realtime subscriptions push ride status changes to rider UI instantly")
add_bullet(tf, "8 database migrations covering riders, drivers, rides, payments, safety")
add_bullet(tf, "Scale target: 30 riders, 15 drivers, 3 admins — kept simple, not over-engineered")

# ─── Slide 4: Demo — Rider Flow ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.5, 11, 0.8, "Live Demo: Rider Flow", size=36, bold=True)
tf = add_text(slide, 0.8, 1.6, 11, 5, "", size=18)
add_bullet(tf, "1. Register / Login  ->  session cookie + riders row created")
add_bullet(tf, "2. Profile  ->  real name, phone, child profiles from Supabase")
add_bullet(tf, "3. Book a Ride  ->  inserts ride with status 'matching'")
add_bullet(tf, "4. Ride Tracking  ->  Supabase Realtime updates status live")
add_bullet(tf, "5. Ride Passes  ->  Stripe checkout in test mode")
add_bullet(tf, "6. Payment Methods  ->  save cards via Stripe Setup Intents")
add_bullet(tf, "7. Flag a Driver  ->  writes to driver_flags table")
add_text(slide, 0.8, 6.0, 11, 0.5, "[ LIVE DEMO ]", size=20, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

# ─── Slide 5: Demo — Driver Flow ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.5, 11, 0.8, "Live Demo: Driver Flow", size=36, bold=True)
tf = add_text(slide, 0.8, 1.6, 11, 5, "", size=18)
add_bullet(tf, "1. Driver Login  ->  middleware redirects to /driver (not rider home)")
add_bullet(tf, "2. Shift Board  ->  real driver name, status, queue count")
add_bullet(tf, "3. Queue  ->  shows real ride requests from riders")
add_bullet(tf, "4. Accept Trip  ->  ride status transitions atomically with version lock")
add_bullet(tf, "5. Pickup Confirmation  ->  two-step identity verification")
add_bullet(tf, "6. Trip Completion  ->  driver back to 'available', rider sees 'completed'")
add_text(slide, 0.8, 6.0, 11, 0.5, "[ LIVE DEMO ]", size=20, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

# ─── Slide 6: Demo — Admin Flow ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.5, 11, 0.8, "Live Demo: Admin Panels", size=36, bold=True)
tf = add_text(slide, 0.8, 1.6, 11, 5, "", size=18)
add_bullet(tf, "1. /admin/drivers  ->  driver fleet with search + filters")
add_bullet(tf, "2. /admin/requests  ->  unfilled ride requests, status/type filters")
add_bullet(tf, "3. /admin/rides  ->  active rides in progress")
add_bullet(tf, "4. /admin/completed  ->  ride history with date range")
add_bullet(tf, "5. /admin/flags  ->  rider complaints — resolve or dismiss")
add_bullet(tf, "6. All tables query real Supabase data with admin-only auth guard")
add_text(slide, 0.8, 6.0, 11, 0.5, "[ LIVE DEMO ]", size=20, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

# ─── Slide 7: Edge Cases & Error Handling ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.5, 11, 0.8, "Edge Cases & Error Handling", size=36, bold=True)
tf = add_text(slide, 0.8, 1.5, 6, 5.5, "Auth & Access", size=22, bold=True, color=ACCENT)
add_bullet(tf, "Unauthenticated user visits /book -> redirected to /login")
add_bullet(tf, "Rider visits /admin -> redirected to rider home")
add_bullet(tf, "Driver visits /passes -> redirected to driver home")
add_bullet(tf, "Expired session -> middleware catches, redirects to /login")
tf2 = add_text(slide, 7, 1.5, 6, 5.5, "Data Integrity", size=22, bold=True, color=GREEN)
add_bullet(tf2, "Double-accept race: optimistic locking via version column")
add_bullet(tf2, "Fare split amount tampering: DB trigger blocks invitee edits")
add_bullet(tf2, "Ride pass double-decrement: version check prevents race")
add_bullet(tf2, "Child profile cross-user: ownership guard checks rider_id + is_child")
add_bullet(tf2, "Expired fare splits: UPDATE checks expires_at > now()")

# ─── Slide 8: Edge Cases (continued) ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.5, 11, 0.8, "Edge Cases (cont.)", size=36, bold=True)
tf = add_text(slide, 0.8, 1.5, 6, 5.5, "Validation & Input", size=22, bold=True, color=ACCENT)
add_bullet(tf, "All server actions validate with Zod before DB queries")
add_bullet(tf, "Empty ride queue -> 'No ride requests' empty state (not mock data)")
add_bullet(tf, "Invalid plan ID -> 'Invalid ride pass plan' error")
add_bullet(tf, "Cancel already-cancelled ride -> rejected with message")
add_bullet(tf, "Cancel already-cancelled pass -> rejected with message")
tf2 = add_text(slide, 7, 1.5, 6, 5.5, "Graceful Degradation", size=22, bold=True, color=GREEN)
add_bullet(tf2, "No Supabase config -> falls back to mock data for dev")
add_bullet(tf2, "DB query error vs not-found: PGRST116 returns null, other errors surface")
add_bullet(tf2, "Profile page works when logged out (shows 'Guest' fallback)")
add_bullet(tf2, "Ride status page uses Realtime + polling fallback every 5s")

# ─── Slide 9: Test Coverage ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.5, 11, 0.8, "Testing", size=36, bold=True)
tf = add_text(slide, 0.8, 1.6, 11, 5, "", size=18)
add_bullet(tf, "Unit tests (Vitest): 190+ tests across all features")
add_bullet(tf, "E2E tests (Playwright): 30+ tests covering full user flows")
add_bullet(tf, "TDD Red-Green-Commit workflow enforced per MASTER_PROMPT")
add_bullet(tf, "Backend suite: server actions tested with mocked Supabase client")
add_bullet(tf, "Integration tests: schema validation against real Supabase")
add_bullet(tf, "Every PR required passing tests before merge")
add_bullet(tf, "")
p = add_bullet(tf, "[ RUN TESTS LIVE ]", size=20, color=ACCENT)
p.font.bold = True

# ─── Slide 10: Reflection Q1 — Jacob ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.3, 11, 0.5, "Reflection", size=16, color=GRAY)
add_text(slide, 0.8, 0.7, 11, 1, "How effective was the LLM in generating backend code?", size=28, bold=True)
add_text(slide, 0.8, 1.6, 11, 0.4, "Jacob", size=20, color=ACCENT, bold=True)
tf = add_text(slide, 0.8, 2.2, 11, 4.5, "", size=18)
add_bullet(tf, "Extremely effective for boilerplate — Supabase queries, Zod schemas,")
add_bullet(tf, "  RLS policies, type mappings all generated correctly first try", level=1, color=GRAY)
add_bullet(tf, "The TDD cycle worked well with Claude — it would write the failing test,")
add_bullet(tf, "  I'd confirm it failed, then it wrote the minimum code to pass", level=1, color=GRAY)
add_bullet(tf, "Server actions, DB migrations, and test mocks were all high quality")
add_bullet(tf, "Liked: consistent patterns — once one feature worked (ride-scheduling),")
add_bullet(tf, "  every subsequent feature followed the same structure", level=1, color=GRAY)
add_bullet(tf, "Generated 8 migrations, 50+ server actions, 190+ tests in ~2 weeks")

# ─── Slide 11: Reflection Q2 — Cassie ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.3, 11, 0.5, "Reflection", size=16, color=GRAY)
add_text(slide, 0.8, 0.7, 11, 1, "What was wrong with what the LLM first generated?", size=28, bold=True)
add_text(slide, 0.8, 1.6, 11, 0.4, "Cassie", size=20, color=ACCENT, bold=True)
tf = add_text(slide, 0.8, 2.2, 11, 4.5, "", size=18)
add_bullet(tf, "Auth model was the biggest issue — actions accepted userId as a parameter")
add_bullet(tf, "  instead of reading from the session. Required a design-level fix.", level=1, color=GRAY)
add_bullet(tf, "Missing dependencies — leaflet, Stripe packages referenced but not installed")
add_bullet(tf, "  Easy to fix (npm install), but broke the build until caught", level=1, color=GRAY)
add_bullet(tf, "Race conditions — fare split accept/decline had TOCTOU bug between")
add_bullet(tf, "  reading status and updating. Fixed with atomic WHERE predicates.", level=1, color=GRAY)
add_bullet(tf, "Harder to fix: getting the LLM to stop using env vars for identity")
add_bullet(tf, "  and use session cookies instead. Took multiple PRs to fully resolve.", level=1, color=GRAY)
add_bullet(tf, "Demo state file (/tmp) cached stale ride statuses — subtle bug")

# ─── Slide 12: Reflection Q3 — Derron ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 0.8, 0.3, 11, 0.5, "Reflection", size=16, color=GRAY)
add_text(slide, 0.8, 0.7, 11, 1.2, "How did you convince yourself the implementation\nwas complete and accomplished your user stories?", size=28, bold=True)
add_text(slide, 0.8, 1.9, 11, 0.4, "Derron", size=20, color=ACCENT, bold=True)
tf = add_text(slide, 0.8, 2.5, 11, 4.5, "", size=18)
add_bullet(tf, "Mapped every user story (US01-US25) to GitHub issues and PRs")
add_bullet(tf, "TDD gave us confidence: every behavior has a test that failed first")
add_bullet(tf, "PR review process: every PR got a code review before merge")
add_bullet(tf, "  LLM helped write reviews — checked for regressions, auth issues, patterns", level=1, color=GRAY)
add_bullet(tf, "End-to-end demo: tonight's flow proves rider->driver->admin works live")
add_bullet(tf, "Admin panels let us verify data in real time — flags, rides, requests")
add_bullet(tf, "Used the LLM to audit test coverage and find gaps")
add_bullet(tf, "  e.g., it flagged missing unit tests in PR #51, missing e2e in PR #50", level=1, color=GRAY)

# ─── Slide 13: Thank You ───
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_text(slide, 1, 2.5, 11, 1.5, "Thank You", size=44, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, 1, 4.0, 11, 0.6, "Ultra — Predictable, Affordable Rides", size=24, color=ACCENT, align=PP_ALIGN.CENTER)
add_text(slide, 1, 5.0, 11, 0.6, "github.com/ai4sd-s26-memphis/team-citrine", size=16, color=GRAY, align=PP_ALIGN.CENTER)

prs.save("meta-documents/ultra-p3-demo.pptx")
print("Saved: meta-documents/ultra-p3-demo.pptx")
print(f"Slides: {len(prs.slides)}")
